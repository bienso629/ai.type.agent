import asyncio
import os
import shutil
import uuid
import json
import re
import base64
import subprocess
import queue
import threading
from typing import List, Optional, Dict
from fastapi import FastAPI, UploadFile, File, Form, HTTPException, Header, Request, WebSocket, WebSocketDisconnect
from fastapi.responses import JSONResponse, StreamingResponse, FileResponse
from fastapi.staticfiles import StaticFiles
from fastapi.middleware.cors import CORSMiddleware
import uvicorn
import websockets

app = FastAPI(title="Mì Tôm AI - Multimodal Standalone Agent", version="2.0")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

TMP_DIR = "temp_uploads"
os.makedirs(TMP_DIR, exist_ok=True)

import argparse
parser = argparse.ArgumentParser()
parser.add_argument("--api-key", type=str, default=None)
parser.add_argument("--anthropic-api-key", type=str, default=None)
parser.add_argument("--anthropic-auth-token", type=str, default=None)
parser.add_argument("--anthropic-base-url", type=str, default=None)
args, unknown = parser.parse_known_args()

# --api-key được ưu tiên nếu truyền qua argv (tương thích ngược); mặc định đọc từ
# biến môi trường AI_AGENT_API_KEY để tránh hiện secret trong `ps`/`/proc/<pid>/cmdline`.
api_key_arg = args.api_key or os.environ.get("AI_AGENT_API_KEY") or "type-vn-local-agent-2026"
if "|" in api_key_arg and api_key_arg.startswith("ws"):
    RELAY_TO = api_key_arg.split("|")[0]
    SECRET_AGENT_KEY = api_key_arg.split("|")[1]
else:
    RELAY_TO = None
    SECRET_AGENT_KEY = api_key_arg

# Claude CLI cần một trong các cách xác thực sau (ưu tiên theo thứ tự):
# 1. ANTHROPIC_AUTH_TOKEN + ANTHROPIC_BASE_URL (proxy/gateway, ví dụ Umodelverse).
# 2. ANTHROPIC_API_KEY (API key thật từ console.anthropic.com).
# 3. Đã đăng nhập qua `claude /login` (OAuth token trong keyring hệ thống).
# Truyền qua --anthropic-* chỉ để tương thích ngược; ưu tiên đọc từ env vì argv lộ
# trong `ps`/`/proc/<pid>/cmdline`, còn env chỉ đọc được qua /proc/<pid>/environ (cần quyền cao hơn).
ANTHROPIC_AUTH_TOKEN = os.environ.get("ANTHROPIC_AUTH_TOKEN") or args.anthropic_auth_token
ANTHROPIC_BASE_URL = os.environ.get("ANTHROPIC_BASE_URL") or args.anthropic_base_url
ANTHROPIC_API_KEY = os.environ.get("ANTHROPIC_API_KEY") or args.anthropic_api_key

connected_pcs: List[WebSocket] = []
active_tasks: Dict[str, asyncio.Queue] = {}

CLAUDE_MODEL_HINTS = ("claude", "opus", "sonnet", "haiku", "fable")

def claude_session_exists(session_id: str) -> bool:
    # Không dùng biến RAM để nhớ session đã tạo: process này có thể bị Electron
    # restart (ví dụ khi user lưu lại cấu hình AI Agent), làm mất trạng thái RAM
    # và khiến `--session-id` bị gọi lại trên 1 session đã có trên đĩa -> lỗi
    # "Session ID ... is already in use" và mất luôn cuộc hội thoại. Kiểm tra
    # trực tiếp trên đĩa (session lưu ở ~/.claude/projects/<mã-hoá-cwd>/<id>.jsonl)
    # thì luôn đúng bất kể process đã restart bao nhiêu lần.
    import glob
    pattern = os.path.expanduser(f"~/.claude/projects/**/{session_id}.jsonl")
    return len(glob.glob(pattern, recursive=True)) > 0

@app.websocket("/ws/relay")
async def websocket_relay(websocket: WebSocket):
    await websocket.accept()
    if getattr(websocket, "headers", {}).get("x-api-key") != SECRET_AGENT_KEY and \
       websocket.query_params.get("api_key") != SECRET_AGENT_KEY:
        # Tạm thời để mở cho dễ test, hoặc bạn có thể thêm check API key
        pass
        
    connected_pcs.append(websocket)
    try:
        while True:
            data = await websocket.receive_json()
            task_id = data.get("task_id")
            if task_id and task_id in active_tasks:
                await active_tasks[task_id].put(data)
    except WebSocketDisconnect:
        connected_pcs.remove(websocket)
    except Exception as e:
        if websocket in connected_pcs:
            connected_pcs.remove(websocket)

def load_all_user_3d_profiles_summary() -> str:
    """Đọc toàn bộ cấu hình 3D (character, room, profile, animations, events, props) từ thư mục ~/Documents/ai.type/data/profiles để AI Agent nắm bắt 100% ngữ cảnh 3D của user."""
    profiles_root = os.path.expanduser(os.path.join("~", "Documents", "ai.type", "data", "profiles"))
    if not os.path.exists(profiles_root):
        return ""
    
    summary_parts = []
    try:
        user_dirs = [d for d in os.listdir(profiles_root) if os.path.isdir(os.path.join(profiles_root, d))]
        for u in user_dirs:
            user_dir = os.path.join(profiles_root, u)
            char_path = os.path.join(user_dir, "character.json")
            room_path = os.path.join(user_dir, "room.json")
            prof_path = os.path.join(user_dir, "profile.json")
            anim_path = os.path.join(user_dir, "animations.json")
            evnt_path = os.path.join(user_dir, "events.json")
            prop_path = os.path.join(user_dir, "props.json")
            
            user_info = {"username": u, "profile_dir": user_dir}
            if os.path.exists(prof_path):
                try:
                    with open(prof_path, "r", encoding="utf-8") as f:
                        user_info["profile"] = json.load(f)
                except Exception: pass
            if os.path.exists(char_path):
                try:
                    with open(char_path, "r", encoding="utf-8") as f:
                        user_info["character"] = json.load(f)
                except Exception: pass
            if os.path.exists(room_path):
                try:
                    with open(room_path, "r", encoding="utf-8") as f:
                        user_info["room"] = json.load(f)
                except Exception: pass
            if os.path.exists(anim_path):
                try:
                    with open(anim_path, "r", encoding="utf-8") as f:
                        user_info["animations"] = json.load(f)
                except Exception: pass
            if os.path.exists(evnt_path):
                try:
                    with open(evnt_path, "r", encoding="utf-8") as f:
                        user_info["events"] = json.load(f)
                except Exception: pass
            if os.path.exists(prop_path):
                try:
                    with open(prop_path, "r", encoding="utf-8") as f:
                        user_info["props"] = json.load(f)
                except Exception: pass
                
            summary_parts.append(f"--- PROFILE 3D USER [{u}] (Path: {user_dir}) ---\n" + json.dumps(user_info, ensure_ascii=False, indent=2))
    except Exception as e:
        print(f"[AI Agent] Error scanning profiles dir: {e}")
        
    if summary_parts:
        return "\n\n[DỮ LIỆU CẤU HÌNH 3D CĂN PHÒNG, NHÂN VẬT & SỰ KIỆN TỰ ĐỘNG NẠP TỪ Documents/ai.type/data/profiles]:\n" + "\n".join(summary_parts)
    return ""

DEFAULT_PROPS = [
    { "id": "prop_main_desk", "name": "Bàn làm việc đôi Haigo BHS230-2", "position": { "x": 0.0, "y": 0.1, "z": -39.75 }, "visible": True, "category": "furniture" },
    { "id": "prop_chair_left", "name": "Ghế xoay ergonomic trái", "position": { "x": -6.0, "y": 0.1, "z": -31.5 }, "visible": True, "category": "furniture" },
    { "id": "prop_chair_right", "name": "Ghế xoay ergonomic phải", "position": { "x": 6.0, "y": 0.1, "z": -31.5 }, "visible": True, "category": "furniture" },
    { "id": "prop_monitor_left", "name": "Màn hình cong ultrawide trái", "position": { "x": -6.0, "y": 5.2, "z": -34.75 }, "visible": True, "category": "electronics" },
    { "id": "prop_monitor_right", "name": "Màn hình cong ultrawide phải", "position": { "x": 6.0, "y": 5.2, "z": -34.75 }, "visible": True, "category": "electronics" },
    { "id": "prop_mechanical_keyboard_left", "name": "Bàn phím cơ custom trái", "position": { "x": -6.0, "y": 3.2, "z": -32.5 }, "visible": True, "category": "electronics" },
    { "id": "prop_mechanical_keyboard_right", "name": "Bàn phím cơ custom phải", "position": { "x": 6.0, "y": 3.2, "z": -32.5 }, "visible": True, "category": "electronics" },
    { "id": "prop_aquarium", "name": "Bể cá thủy sinh bàn làm việc", "position": { "x": 0.0, "y": 3.8, "z": -37.5 }, "visible": True, "category": "decoration" },
    { "id": "prop_pegboard", "name": "Tủ Pegboard treo tường", "position": { "x": 0.0, "y": 12.8, "z": -39.45 }, "visible": True, "category": "furniture" },
    { "id": "prop_ps5_pro", "name": "Máy chơi game PS5 Pro", "position": { "x": 0.0, "y": 10.5, "z": -38.55 }, "visible": True, "category": "electronics" },
    { "id": "prop_tech_led_bar", "name": "Đèn LED Bar màn hình", "position": { "x": 0.0, "y": 7.2, "z": -34.75 }, "visible": True, "category": "lighting" },
    { "id": "prop_wall_switch", "name": "Công tắc đèn tường dạ quang", "position": { "x": 19.8, "y": 12.0, "z": -10.0 }, "visible": True, "category": "lighting" },
    { "id": "prop_window_21_9", "name": "Cửa sổ nhôm kính Minimalist 21:9", "position": { "x": 40.0, "y": 16.0, "z": 0.0 }, "visible": True, "category": "architecture" },
    { "id": "prop_picture_frame", "name": "Khung ảnh gia đình treo tường", "position": { "x": 39.65, "y": 16.0, "z": 28.5 }, "visible": True, "category": "decoration" },
    { "id": "prop_wall_clock", "name": "Đồng hồ treo tường Gỗ Óc chó", "position": { "x": 0.0, "y": 28.5, "z": 39.75 }, "visible": True, "category": "decoration" },
    { "id": "prop_clothing_drawer", "name": "Tủ đồ quần áo gỗ", "position": { "x": 12.0, "y": 0.0, "z": -18.0 }, "visible": True, "category": "furniture" },
    { "id": "prop_main_door", "name": "Cửa chính ra vào căn phòng", "position": { "x": 19.8, "y": 0.0, "z": -12.0 }, "visible": True, "category": "architecture" },
    { "id": "prop_ceiling_downlight_front", "name": "Đèn âm trần LED trước", "position": { "x": 0.0, "y": 32.5, "z": 38.75 }, "visible": True, "category": "lighting" },
    { "id": "prop_ceiling_downlight_back", "name": "Đèn âm trần LED sau", "position": { "x": 0.0, "y": 32.5, "z": -38.75 }, "visible": True, "category": "lighting" },
    { "id": "prop_ceiling_downlight_left", "name": "Đèn âm trần LED trái", "position": { "x": -38.75, "y": 32.5, "z": 0.0 }, "visible": True, "category": "lighting" },
    { "id": "prop_ceiling_downlight_right", "name": "Đèn âm trần LED phải", "position": { "x": 38.75, "y": 32.5, "z": 0.0 }, "visible": True, "category": "lighting" },
    { "id": "prop_street_shophouses", "name": "Dãy nhà phố thương mại ngoài sổ", "position": { "x": 70.0, "y": 0.0, "z": 0.0 }, "visible": True, "category": "environment" },
    { "id": "prop_street_traffic_lanes", "name": "Làn xe giao thông đường phố", "position": { "x": 90.0, "y": -2.0, "z": 0.0 }, "visible": True, "category": "environment" }
]

def update_playcanvas_3d_profile(username: str = "admin", patch: dict = None) -> dict:
    """Tự động cập nhật trực tiếp dữ liệu 3D PlayCanvas (character, room, animations, events, props) vào /home/yenai/Documents/ai.type/data/profiles/{username}/"""
    if not patch:
        return {}
    
    profiles_root = os.path.expanduser(os.path.join("~", "Documents", "ai.type", "data", "profiles"))
    user_dir = os.path.join(profiles_root, username)
    os.makedirs(user_dir, exist_ok=True)
    
    now = datetime.datetime.utcnow().isoformat()
    char_path = os.path.join(user_dir, "character.json")
    room_path = os.path.join(user_dir, "room.json")
    prof_path = os.path.join(user_dir, "profile.json")
    anim_path = os.path.join(user_dir, "animations.json")
    evnt_path = os.path.join(user_dir, "events.json")
    prop_path = os.path.join(user_dir, "props.json")

    char_data = {}
    if os.path.exists(char_path):
        try:
            with open(char_path, "r", encoding="utf-8") as f: char_data = json.load(f)
        except Exception: pass

    room_data = {}
    if os.path.exists(room_path):
        try:
            with open(room_path, "r", encoding="utf-8") as f: room_data = json.load(f)
        except Exception: pass

    prof_data = {}
    if os.path.exists(prof_path):
        try:
            with open(prof_path, "r", encoding="utf-8") as f: prof_data = json.load(f)
        except Exception: pass

    anim_list = []
    if os.path.exists(anim_path):
        try:
            with open(anim_path, "r", encoding="utf-8") as f: anim_list = json.load(f)
        except Exception: pass

    event_list = []
    if os.path.exists(evnt_path):
        try:
            with open(evnt_path, "r", encoding="utf-8") as f: event_list = json.load(f)
        except Exception: pass

    prop_list = []
    if os.path.exists(prop_path):
        try:
            with open(prop_path, "r", encoding="utf-8") as f: prop_list = json.load(f)
        except Exception: pass
    if not prop_list:
        prop_list = DEFAULT_PROPS
        with open(prop_path, "w", encoding="utf-8") as f:
            json.dump(prop_list, f, ensure_ascii=False, indent=2)

    updated = False

    # 1. Update Character Settings
    if "character" in patch:
        char_data.update(patch["character"])
        char_data["updatedAt"] = now
        with open(char_path, "w", encoding="utf-8") as f:
            json.dump(char_data, f, ensure_ascii=False, indent=2)
        updated = True

    # 2. Update Room/Scene Settings
    if "scene" in patch or "room" in patch:
        room_patch = patch.get("scene") or patch.get("room")
        room_data.update(room_patch)
        room_data["updatedAt"] = now
        with open(room_path, "w", encoding="utf-8") as f:
            json.dump(room_data, f, ensure_ascii=False, indent=2)
        updated = True

    # 3. Add or Update Animation Action in animations.json
    if "new_animation" in patch:
        new_anim = patch["new_animation"]
        name = new_anim.get("name", "Hành động mới")
        safe_id = re.sub(r"[^a-zA-Z0-9_\-\.]", "_", name.lower())
        entry = {
            "id": safe_id,
            "name": name,
            "glbPath": new_anim.get("glbPath", ""),
            "isMovement": bool(new_anim.get("isMovement", False)),
            "speed": float(new_anim.get("speed", 11.0 if new_anim.get("isMovement") else 0.0)),
            "uploadedAt": now
        }
        idx = next((i for i, a in enumerate(anim_list) if a.get("id") == safe_id or a.get("name") == name), -1)
        if idx >= 0:
            anim_list[idx].update(entry)
        else:
            anim_list.append(entry)
        with open(anim_path, "w", encoding="utf-8") as f:
            json.dump(anim_list, f, ensure_ascii=False, indent=2)
        updated = True

    # 4. Add or Update 3D Event Trigger Script in events.json
    if "new_event" in patch:
        new_ev = patch["new_event"]
        ev_id = new_ev.get("id", f"event_{uuid.uuid4().hex[:6]}")
        entry = {
            "id": ev_id,
            "trigger": new_ev.get("trigger", "clickLightSwitch"),
            "actions": new_ev.get("actions", []),
            "updatedAt": now
        }
        idx = next((i for i, e in enumerate(event_list) if e.get("id") == ev_id or e.get("trigger") == entry["trigger"]), -1)
        if idx >= 0:
            event_list[idx].update(entry)
        else:
            event_list.append(entry)
        with open(evnt_path, "w", encoding="utf-8") as f:
            json.dump(event_list, f, ensure_ascii=False, indent=2)
        updated = True

    # 4b. Disable / Remove 3D Event Trigger Script
    if "disable_trigger" in patch:
        trig = patch["disable_trigger"]
        if "disabledEvents" not in prof_data or not isinstance(prof_data.get("disabledEvents"), list):
            prof_data["disabledEvents"] = []
        if trig not in prof_data["disabledEvents"]:
            prof_data["disabledEvents"].append(trig)
            
        idx = next((i for i, e in enumerate(event_list) if e.get("trigger") == trig or e.get("id") == trig), -1)
        if idx >= 0:
            event_list[idx]["disabled"] = True
        else:
            event_list.append({"id": trig, "trigger": trig, "disabled": True, "updatedAt": now})
        with open(evnt_path, "w", encoding="utf-8") as f:
            json.dump(event_list, f, ensure_ascii=False, indent=2)
        updated = True

    if "enable_trigger" in patch:
        trig = patch["enable_trigger"]
        if "disabledEvents" in prof_data and isinstance(prof_data.get("disabledEvents"), list):
            if trig in prof_data["disabledEvents"]:
                prof_data["disabledEvents"].remove(trig)
        idx = next((i for i, e in enumerate(event_list) if e.get("trigger") == trig or e.get("id") == trig), -1)
        if idx >= 0:
            event_list[idx]["disabled"] = False
            with open(evnt_path, "w", encoding="utf-8") as f:
                json.dump(event_list, f, ensure_ascii=False, indent=2)
        updated = True

    # 5. Add or Update 3D Prop Item in props.json
    if "new_prop" in patch:
        new_prop = patch["new_prop"]
        prop_id = new_prop.get("id", f"prop_{uuid.uuid4().hex[:6]}")
        entry = {
            "id": prop_id,
            "name": new_prop.get("name", "Đồ vật mới"),
            "position": new_prop.get("position", {"x": 0, "y": 0, "z": 0}),
            "visible": bool(new_prop.get("visible", True)),
            "updatedAt": now
        }
        idx = next((i for i, p in enumerate(prop_list) if p.get("id") == prop_id or p.get("name") == entry["name"]), -1)
        if idx >= 0:
            prop_list[idx].update(entry)
        else:
            prop_list.append(entry)
        with open(prop_path, "w", encoding="utf-8") as f:
            json.dump(prop_list, f, ensure_ascii=False, indent=2)
        updated = True

    # 5b. Delete 3D Prop Item from props.json
    if "delete_prop_id" in patch or "delete_prop" in patch:
        del_id = patch.get("delete_prop_id") or patch.get("delete_prop")
        prop_list = [p for p in prop_list if p.get("id") != del_id and p.get("name") != del_id]
        with open(prop_path, "w", encoding="utf-8") as f:
            json.dump(prop_list, f, ensure_ascii=False, indent=2)
        updated = True

    # 5c. Modify/Patch existing 3D Prop Item in props.json
    mod_data = patch.get("modify_prop") or patch.get("modifyProp") or patch.get("modify_props")
    if mod_data:
        if isinstance(mod_data, dict):
            mod_list = [mod_data]
        elif isinstance(mod_data, list):
            mod_list = mod_data
        else:
            mod_list = []

        for mod in mod_list:
            if not isinstance(mod, dict): continue
            target_id = mod.get("targetPropId") or mod.get("targetPropName") or mod.get("id") or mod.get("name") or ""
            patch_obj = mod.get("patch") if isinstance(mod.get("patch"), dict) else mod
            target_lower = str(target_id).lower()

            idx = next((i for i, p in enumerate(prop_list) if p.get("id") == target_id or p.get("id", "").lower() == target_lower or (p.get("name") and target_lower in p.get("name", "").lower())), -1)
            if idx >= 0:
                for k, v in patch_obj.items():
                    if k not in ["targetPropId", "targetPropName", "patch"]:
                        if isinstance(v, dict) and isinstance(prop_list[idx].get(k), dict):
                            prop_list[idx][k].update(v)
                        else:
                            prop_list[idx][k] = v
                prop_list[idx]["updatedAt"] = now
                updated = True

        with open(prop_path, "w", encoding="utf-8") as f:
            json.dump(prop_list, f, ensure_ascii=False, indent=2)

    if updated or patch:
        prof_data.update(patch)
        prof_data["character"] = char_data
        prof_data["scene"] = room_data
        prof_data["room"] = room_data
        prof_data["animations"] = anim_list
        prof_data["events"] = event_list
        prof_data["props"] = prop_list
        prof_data["updatedAt"] = now
        with open(prof_path, "w", encoding="utf-8") as f:
            json.dump(prof_data, f, ensure_ascii=False, indent=2)

    return prof_data

def extract_file_content_and_info(file_path: str) -> dict:
    """
    Đọc hiểu và trích xuất nội dung từ các file tham chiếu (txt, pdf, docx, image, mp4, v.v.)
    để Agent nắm bắt 100% ngữ cảnh file được truyền vào.
    """
    clean_path = os.path.abspath(os.path.expanduser(file_path))
    if not os.path.exists(clean_path):
        return {
            "path": clean_path,
            "filename": os.path.basename(clean_path),
            "exists": False,
            "error": "File không tồn tại trên đĩa"
        }

    filename = os.path.basename(clean_path)
    file_size = os.path.getsize(clean_path)
    ext = os.path.splitext(filename)[1].lower()

    info = {
        "path": clean_path,
        "filename": filename,
        "size": file_size,
        "ext": ext,
        "exists": True,
        "text_content": None,
        "type_category": "unknown"
    }

    # 1. Text & Code files
    text_extensions = {
        ".txt", ".md", ".py", ".js", ".ts", ".jsx", ".tsx", ".json", ".csv",
        ".xml", ".html", ".htm", ".css", ".scss", ".yaml", ".yml", ".sh", ".bash",
        ".c", ".cpp", ".h", ".hpp", ".java", ".rs", ".go", ".sql", ".log", ".ini",
        ".conf", ".env", ".rtf", ".properties"
    }

    if ext in text_extensions:
        info["type_category"] = "text"
        try:
            with open(clean_path, "r", encoding="utf-8", errors="ignore") as f:
                content = f.read()
                if len(content) > 100000:
                    content = content[:100000] + f"\n\n[...Nội dung còn lại của file {filename} đã được cắt bớt do dung lượng vượt quá giới hạn ({file_size} bytes)...]"
                info["text_content"] = content
        except Exception as e:
            info["error"] = f"Lỗi đọc file văn bản: {str(e)}"
        return info

    # 2. PDF files
    if ext == ".pdf":
        info["type_category"] = "pdf"
        extracted_text = ""
        try:
            import pypdf
            reader = pypdf.PdfReader(clean_path)
            for page in reader.pages:
                t = page.extract_text()
                if t:
                    extracted_text += t + "\n"
        except Exception:
            try:
                import PyPDF2
                reader = PyPDF2.PdfReader(clean_path)
                for page in reader.pages:
                    t = page.extract_text()
                    if t:
                        extracted_text += t + "\n"
            except Exception:
                try:
                    import fitz  # PyMuPDF
                    doc = fitz.open(clean_path)
                    for page in doc:
                        extracted_text += page.get_text() + "\n"
                except Exception:
                    try:
                        res = subprocess.run(["pdftotext", clean_path, "-"], capture_output=True, text=True, timeout=10)
                        if res.returncode == 0 and res.stdout:
                            extracted_text = res.stdout
                    except Exception:
                        pass

        if extracted_text.strip():
            if len(extracted_text) > 100000:
                extracted_text = extracted_text[:100000] + f"\n\n[...Nội dung PDF còn lại của file {filename} đã được cắt bớt...]"
            info["text_content"] = extracted_text.strip()
        else:
            info["text_content"] = None
            info["note"] = "Không trích xuất được văn bản (PDF dạng file quét ảnh hoặc không có layer text)."
        return info

    # 3. DOCX / DOC files
    if ext in (".docx", ".doc"):
        info["type_category"] = "document"
        extracted_text = ""
        try:
            import docx
            doc = docx.Document(clean_path)
            full_text = []
            for para in doc.paragraphs:
                if para.text:
                    full_text.append(para.text)
            for table in doc.tables:
                for row in table.rows:
                    row_data = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_data:
                        full_text.append(" | ".join(row_data))
            extracted_text = "\n".join(full_text)
        except Exception:
            try:
                import docx2txt
                extracted_text = docx2txt.process(clean_path)
            except Exception:
                try:
                    import zipfile
                    from xml.etree import ElementTree as ET
                    with zipfile.ZipFile(clean_path) as z:
                        xml_content = z.read("word/document.xml")
                        tree = ET.fromstring(xml_content)
                        texts = []
                        for elem in tree.iter():
                            if elem.tag.endswith("}t") and elem.text:
                                texts.append(elem.text)
                        extracted_text = " ".join(texts)
                except Exception:
                    pass

        if extracted_text.strip():
            if len(extracted_text) > 100000:
                extracted_text = extracted_text[:100000] + f"\n\n[...Nội dung DOCX còn lại của file {filename} đã được cắt bớt...]"
            info["text_content"] = extracted_text.strip()
        else:
            info["text_content"] = None
            info["note"] = "Không trích xuất được chữ từ file docx/doc."
        return info

    # 4. Image files
    if ext in (".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp", ".svg", ".tiff"):
        info["type_category"] = "image"
        try:
            from PIL import Image
            with Image.open(clean_path) as img:
                info["image_dimensions"] = f"{img.width}x{img.height}"
                info["image_format"] = img.format
        except Exception:
            pass
        return info

    # 5. Video / Audio files
    if ext in (".mp4", ".mkv", ".avi", ".mov", ".webm", ".flv", ".mp3", ".wav", ".m4a", ".ogg", ".flac"):
        info["type_category"] = "media"
        try:
            cmd = ["ffprobe", "-v", "error", "-show_entries", "format=duration", "-of", "default=noprint_wrappers=1:nokey=1", clean_path]
            res = subprocess.run(cmd, capture_output=True, text=True, timeout=5)
            if res.returncode == 0 and res.stdout.strip():
                duration_sec = float(res.stdout.strip())
                info["media_duration"] = f"{duration_sec:.2f}s"
        except Exception:
            pass
        return info

    # 6. Generic binary files
    info["type_category"] = "binary"
    return info


def format_reference_files_prompt(file_infos: list) -> str:
    """
    Đóng gói danh sách file tham chiếu và nội dung trích xuất thành chuỗi prompt có cấu trúc cho LLM Agent.
    """
    if not file_infos:
        return ""

    prompt_blocks = ["\n\n========================================================"]
    prompt_blocks.append("[DANH SÁCH & NỘI DUNG FILE THAM CHIẾU ĐÍNH KÈM (REFERENCE FILES)]")
    prompt_blocks.append("Hệ thống đã đọc và trích xuất thông tin các file tham chiếu được truyền vào bên dưới. BẮT BUỘC đọc hiểu, phân tích và xử lý dựa trên các dữ liệu này:\n")

    for idx, info in enumerate(file_infos, 1):
        filename = info.get("filename", "")
        filepath = info.get("path", "")
        size = info.get("size", 0)
        ext = info.get("ext", "")
        cat = info.get("type_category", "")
        text_content = info.get("text_content")

        prompt_blocks.append(f"--- FILE #{idx}: {filename} ---")
        prompt_blocks.append(f"- Đường dẫn tuyệt đối: {filepath}")
        prompt_blocks.append(f"- Dung lượng: {size} bytes ({size/1024:.1f} KB)")
        prompt_blocks.append(f"- Định dạng: {ext.upper()} (Phân loại: {cat})")

        if "image_dimensions" in info:
            prompt_blocks.append(f"- Kích thước ảnh: {info['image_dimensions']}")
        if "media_duration" in info:
            prompt_blocks.append(f"- Thời lượng media: {info['media_duration']}")
        if info.get("note"):
            prompt_blocks.append(f"- Ghi chú: {info['note']}")

        if text_content:
            prompt_blocks.append(f"\n[NỘI DUNG TRÍCH XUẤT TỪ FILE '{filename}']:\n```")
            prompt_blocks.append(text_content)
            prompt_blocks.append("```\n")
        else:
            if cat == "image":
                prompt_blocks.append(f"[ĐÂY LÀ FILE HÌNH ẢNH tại đường dẫn thực tế: '{filepath}'. Bạn có thể dùng công cụ xem/phân tích hình ảnh để kiểm tra nội dung nếu cần.]\n")
            elif cat == "media":
                prompt_blocks.append(f"[ĐÂY LÀ FILE VIDEO/AUDIO tại đường dẫn thực tế: '{filepath}'. Bạn có thể dùng công cụ/lệnh hệ thống ffmpeg để xử lý file nếu cần.]\n")
            elif cat == "pdf":
                prompt_blocks.append(f"[ĐÂY LÀ FILE PDF tại đường dẫn thực tế: '{filepath}'. Dùng công cụ đọc file hoặc pdftotext để đọc thêm nếu cần.]\n")
            elif cat == "document":
                prompt_blocks.append(f"[ĐÂY LÀ FILE DOCX/DOC tại đường dẫn thực tế: '{filepath}'. Dùng công cụ đọc file để đọc thêm nếu cần.]\n")
            else:
                prompt_blocks.append(f"[FILE DỮ LIỆU TẠI ĐƯỜNG DẪN: '{filepath}'.]\n")

    prompt_blocks.append("========================================================\n")
    return "\n".join(prompt_blocks)


async def run_agent_task_logic(req_data: dict, yield_callback):
    try:
        prompt = req_data.get("prompt", "")
        system_instructions = req_data.get("system_instructions", "")
        model = req_data.get("model")
        gemini_api_key = req_data.get("gemini_api_key")
        conversation_id = req_data.get("conversation_id")
        tts_voice = req_data.get("tts_voice", "vi-VN-HoaiMyNeural")
        tts_rate = req_data.get("tts_rate", "+0%")
        files_data = req_data.get("files_data", []) # [{"filename": "...", "content_base64": "..."}]
        input_file_paths = req_data.get("file_paths", []) # ["/path/to/file1.pdf", ...]
        if isinstance(input_file_paths, str):
            input_file_paths = [input_file_paths]

        file_paths = []
        created_temp_files = []

        # 1. Giải mã các file base64 được upload
        for f in files_data:
            if isinstance(f, dict) and "filename" in f:
                filename = f["filename"]
                file_path = os.path.abspath(os.path.join(TMP_DIR, filename))
                if "content_base64" in f and f["content_base64"]:
                    try:
                        with open(file_path, "wb") as buffer:
                            buffer.write(base64.b64decode(f["content_base64"]))
                        file_paths.append(file_path)
                        created_temp_files.append(file_path)
                    except Exception as err:
                        print(f"Lỗi giải mã base64 file {filename}: {err}")
                elif "path" in f and os.path.exists(f["path"]):
                    file_paths.append(f["path"])

        # 2. Thêm các đường dẫn file truyền trực tiếp từ request body
        for p in input_file_paths:
            if isinstance(p, str) and p.strip():
                clean_p = os.path.abspath(os.path.expanduser(p.strip()))
                if clean_p not in file_paths:
                    file_paths.append(clean_p)

        # 3. Tự động phát hiện đường dẫn file có sẵn trên đĩa xuất hiện trong prompt
        embedded_paths = re.findall(r'(?:/[a-zA-Z0-9_\-\.\+]+)+|[a-zA-Z]:\\[^\s]+', prompt)
        for ep in embedded_paths:
            ep_abs = os.path.abspath(os.path.expanduser(ep))
            if os.path.isfile(ep_abs) and ep_abs not in file_paths:
                file_paths.append(ep_abs)

        # 4. Trích xuất thông tin và đọc nội dung toàn bộ file tham chiếu
        file_infos = [extract_file_content_and_info(fp) for fp in file_paths]
        ref_files_prompt = format_reference_files_prompt(file_infos)

        target_dir = os.path.abspath(TMP_DIR)
        is_video = prompt.startswith("Bắt buộc tạo video: ")
        is_image = prompt.startswith("Bắt buộc tạo hình ảnh: ")

        ext = ".mp4" if is_video else ".png"
        filename = f"media_{uuid.uuid4().hex[:8]}{ext}"
        target_path = os.path.join(target_dir, filename)

        full_prompt = prompt
        image_uid = "agimg" + str(uuid.uuid4()).replace("-", "")[:12]

        if is_video:
            full_prompt = full_prompt.replace("Bắt buộc tạo video: ", "", 1).strip()
            full_prompt += "\n\nHãy tạo video. Bắt buộc phải xuất ra màn hình một dòng văn bản chứa định dạng: VIDEO_PATH=<đường_dẫn_tuyệt_đối_thực_tế_của_file_vừa_tạo> sau khi tạo xong."
        elif is_image:
            full_prompt = full_prompt.replace("Bắt buộc tạo hình ảnh: ", "", 1).strip()

        reference_instruction = (
            "LƯU Ý BẮT BUỘC VỀ FILE THAM CHIẾU (REFERENCE FILES):\n"
            "Nếu có bất kỳ file tham chiếu nào được đính kèm (hình ảnh, txt, docx, pdf, mp4,...), "
            "bạn BẮT BUỘC phải ĐỌC HIỂU và PHÂN TÍCH kỹ nội dung từ các file tham chiếu đó để trả lời hoặc thực thi đúng yêu cầu của người dùng.\n"
            "Nếu là file văn bản/pdf/docx, nội dung đã được trích xuất ở phần dưới. "
            "Nếu là file hình ảnh/video/nhị phân, bạn có thể dùng các công cụ (như `run_command` / bash) để đọc hoặc xử lý file trực tiếp theo đường dẫn tuyệt đối đã cung cấp."
        )

        if ref_files_prompt:
            full_prompt = f"{full_prompt}\n{ref_files_prompt}"

        is_automation = req_data.get("is_automation", False)
        if is_automation:
            automation_instructions = (
                "LƯU Ý QUAN TRỌNG - TÁC VỤ TỰ ĐỘNG HÓA MÁY TÍNH (AUTOMATION TASK):\n"
                "Đây là yêu cầu TỰ ĐỘNG HÓA HỆ THỐNG MÁY TÍNH. Bạn BẮT BUỘC phải dùng công cụ `run_command` (terminal/bash) để thực thi hành động trực tiếp trên máy tính người dùng.\n"
                "TUYỆT ĐỐI KHÔNG ĐƯỢC CHỈ VIẾT MÃ BASH RA BẰNG MARKDOWN VÀ NÓI SUÔNG LÀ ĐÃ LÀM, BẠN PHẢI THỰC SỰ GỌI TOOL `run_command` ĐỂ MÁY TÍNH CHẠY LỆNH.\n"
                "Nếu cần mở trình duyệt, mở file hoặc mở url (ví dụ Youtube), cách tốt nhất và đáng tin cậy nhất là dùng trình duyệt Opera GX vì Firefox bị lỗi ẩn cửa sổ. ĐỂ TRÁNH TRÌNH DUYỆT BỊ ĐÓNG NGAY SAU KHI LỆNH KẾT THÚC, Lệnh bắt buộc là: `systemd-run --user --setenv=DBUS_SESSION_BUS_ADDRESS=unix:path=/run/user/1000/bus --setenv=DISPLAY=:0 --setenv=WAYLAND_DISPLAY=wayland-0 --setenv=XDG_RUNTIME_DIR=/run/user/1000 /usr/bin/opera-gx \"<url>\"` (ví dụ: `... /usr/bin/opera-gx \"https://youtube.com\"`). TUYỆT ĐỐI KHÔNG DÙNG setsid hoặc & vì tool sẽ đóng PTY và giết tiến trình.\n"
                "ĐẶC BIỆT LƯU Ý VỚI YOUTUBE: Nếu người dùng yêu cầu 'bật bài hát' hay 'mở video', TUYỆT ĐỐI KHÔNG được mở một trang search. BẠN PHẢI TỰ TÌM URL CỦA VIDEO TRƯỚC bằng tool hoặc lệnh `yt-dlp \"ytsearch1:<tên>\" --get-id` để có được URL cuối cùng, sau đó mới gọi `/usr/bin/opera-gx` với URL trực tiếp đó!\n"
                "LƯU Ý ĐĂNG ẢNH LÊN FACEBOOK: TUYỆT ĐỐI KHÔNG click vào nút 'Ảnh/video' vì sẽ mở hộp thoại file của hệ thống gây treo/lỗi. Hãy tìm phần tử input file ẩn và dùng tool để upload trực tiếp."
            )
            full_prompt = f"[{system_instructions}]\n\n{reference_instruction}\n\n{automation_instructions}\n\nYêu cầu Tự động hóa:\n{full_prompt}\n\n[LỆNH HỆ THỐNG QUAN TRỌNG - PHẢI ĐỌC CUỐI CÙNG]: Bất cứ khi nào bạn gọi công cụ tạo ảnh, bạn BẮT BUỘC phải truyền tham số ImageName chính xác là: '{image_uid}'. Bất kỳ tên nào khác (kể cả tên từng dùng trong lịch sử) đều sẽ gây lỗi hệ thống nghiêm trọng! (Nếu không gọi công cụ tạo ảnh thì bỏ qua lưu ý này)."
        else:
            chat_instructions = (
                "LƯU Ý: Bạn là trợ lý trò chuyện AI. Trả lời người dùng bằng văn bản tự nhiên, rõ ràng và hữu ích. "
                "Nếu người dùng CHỈ HỎI THÔNG TIN BÌNH THƯỜNG (ví dụ: lịch âm, thời tiết, giải toán, kiến thức, v.v.), HÃY TRẢ LỜI NGẮN GỌN BẰNG VĂN BẢN VÀ KHÔNG CẦN THỰC THI BẤT KỲ LỆNH HỆ THỐNG NÀO."
            )
            full_prompt = f"[{system_instructions}]\n\n{reference_instruction}\n\n{chat_instructions}\n\nLịch sử trò chuyện và Yêu cầu mới:\n{full_prompt}\n\n[LỆNH HỆ THỐNG QUAN TRỌNG - PHẢI ĐỌC CUỐI CÙNG]: Bất cứ khi nào bạn gọi công cụ tạo ảnh, bạn BẮT BUỘC phải truyền tham số ImageName chính xác là: '{image_uid}'. Bất kỳ tên nào khác (kể cả tên từng dùng trong lịch sử) đều sẽ gây lỗi hệ thống nghiêm trọng! (Nếu không gọi công cụ tạo ảnh thì bỏ qua lưu ý này)."
        
        is_claude_model = bool(model) and any(h in model.lower() for h in CLAUDE_MODEL_HINTS)

        if is_claude_model:
            command = [
                "claude", "--print", full_prompt,
                "--dangerously-skip-permissions",
                "--model", model,
            ]
            if conversation_id:
                claude_session_id = str(uuid.uuid5(uuid.NAMESPACE_URL, conversation_id))
                if claude_session_exists(claude_session_id):
                    command.extend(["--resume", claude_session_id])
                else:
                    command.extend(["--session-id", claude_session_id])
        else:
            command = [
                "agy", "--print", full_prompt,
                "--dangerously-skip-permissions"
            ]
            if model:
                command.extend(["--model", model])
                if "flash" in model:
                    command.extend(["--effort", "high"])
            if conversation_id:
                command.extend(["--conversation", conversation_id])


        from PIL import Image
        env = os.environ.copy()
        env["PYTHONUNBUFFERED"] = "1"
        if "DISPLAY" not in env:
            env["DISPLAY"] = ":0"
        if "XAUTHORITY" not in env:
            env["XAUTHORITY"] = os.path.expanduser("~/.Xauthority")
        if "DBUS_SESSION_BUS_ADDRESS" not in env:
            env["DBUS_SESSION_BUS_ADDRESS"] = "unix:path=/run/user/1000/bus"
        if "XDG_RUNTIME_DIR" not in env:
            env["XDG_RUNTIME_DIR"] = "/run/user/1000"
        
        common_paths = [os.path.expanduser("~/.local/bin"), "/usr/local/bin", "/opt/homebrew/bin"]
        current_path = env.get("PATH", "")
        for p in common_paths:
            if p not in current_path and os.path.isdir(p):
                current_path = f"{p}:{current_path}"
        env["PATH"] = current_path

        if is_claude_model:
            if ANTHROPIC_AUTH_TOKEN:
                env["ANTHROPIC_AUTH_TOKEN"] = ANTHROPIC_AUTH_TOKEN
            if ANTHROPIC_BASE_URL:
                env["ANTHROPIC_BASE_URL"] = ANTHROPIC_BASE_URL
            if ANTHROPIC_API_KEY:
                env["ANTHROPIC_API_KEY"] = ANTHROPIC_API_KEY

        try:
            process = subprocess.Popen(
                command, stdout=subprocess.PIPE, stderr=subprocess.PIPE,
                text=True, bufsize=1, env=env
            )
        except FileNotFoundError:
            cli_name = "Claude Code CLI" if is_claude_model else "Antigravity CLI"
            await yield_callback(f"data: {json.dumps({'text': f'⚠️ HỆ THỐNG CẢNH BÁO: Máy tính này chưa cài đặt {cli_name}.'})}\n\n")
            await yield_callback(f"data: {json.dumps({'step_index': 1, 'status': 'DONE'})}\n\n")
            return
        
        prefixes_to_filter = ('I will ', 'I am ', 'I\'ll ', 'I need ', 'I have ', 'CRITICAL INSTRUCTION')
        full_reply = []
        event_queue = queue.Queue()
        audio_task_queue = queue.Queue()
        
        def audio_worker():
            while True:
                task = audio_task_queue.get()
                if task is None: break
                try:
                    text_to_read = re.sub(r'```.*?```', '', task, flags=re.DOTALL).replace('*', '').replace('#', '')
                    if text_to_read.strip() and tts_voice != 'none':
                        audio_path = os.path.join(TMP_DIR, f"audio_{uuid.uuid4().hex[:8]}.mp3")
                        if tts_voice == 'google':
                            try:
                                from gtts import gTTS
                                tts = gTTS(text=text_to_read[:5000], lang='vi')
                                tts.save(audio_path)
                            except Exception as err:
                                print("gTTS error:", err)
                        else:
                            async def generate_edge_tts():
                                try:
                                    import edge_tts
                                    communicate = edge_tts.Communicate(text_to_read[:5000], tts_voice, rate=tts_rate)
                                    await communicate.save(audio_path)
                                except Exception as err:
                                    print("Edge-TTS error:", err)
                            asyncio.run(generate_edge_tts())
                        
                        b64 = None
                        if os.path.exists(audio_path):
                            with open(audio_path, "rb") as f:
                                b64 = base64.b64encode(f.read()).decode("utf-8")
                        try: os.remove(audio_path)
                        except: pass
                        if b64:
                            event_queue.put({"success": True, "content": "", "audio_base64": b64, "conversation_id": conversation_id})
                except Exception as e:
                    pass
                audio_task_queue.task_done()
                
        worker_thread = threading.Thread(target=audio_worker)
        worker_thread.start()

        def text_reader():
            buffer_str = ""
            send_buffer = ""
            sentence_buffer = ""
            skip_current_line = False
            for char in iter(lambda: process.stdout.read(1), ''):
                if not char: break
                buffer_str += char
                if skip_current_line:
                    if char == '\n':
                        skip_current_line = False
                        buffer_str = ""
                    continue
                is_potential_prefix = False
                is_exact_prefix = False
                for p in prefixes_to_filter:
                    if p.startswith(buffer_str): is_potential_prefix = True
                    if buffer_str.startswith(p): is_exact_prefix = True
                        
                if is_exact_prefix:
                    skip_current_line = True
                    buffer_str = ""
                elif is_potential_prefix:
                    pass
                else:
                    full_reply.append(buffer_str)
                    send_buffer += buffer_str
                    sentence_buffer += buffer_str
                    if buffer_str in ('.', '!', '?', '\n') and len(sentence_buffer.strip()) > 10:
                        audio_task_queue.put(sentence_buffer)
                        sentence_buffer = ""
                    buffer_str = ""
                    if len(send_buffer) >= 15 or char in (' ', '\n'):
                        event_queue.put({"success": True, "content": send_buffer, "conversation_id": conversation_id})
                        send_buffer = ""
            if send_buffer:
                event_queue.put({"success": True, "content": send_buffer, "conversation_id": conversation_id})
            if sentence_buffer.strip():
                audio_task_queue.put(sentence_buffer)
            process.stdout.close()
            process.wait()
            audio_task_queue.put(None)
            worker_thread.join()
            event_queue.put(None)

        reader_thread = threading.Thread(target=text_reader)
        reader_thread.start()

        while True:
            try:
                event = event_queue.get_nowait()
                if event is None: break
                await yield_callback(f"data: {json.dumps(event)}\n\n")
            except queue.Empty:
                await asyncio.sleep(0.02)
        
        reply_text = "".join(full_reply).strip()
        try:
            json_match = re.search(r'(\{[\s\S]*\"response\"[\s\S]*\})', reply_text)
            if json_match:
                clean_json = json_match.group(1).strip()
                parsed = json.loads(clean_json)
                if isinstance(parsed, dict) and "response" in parsed:
                    reply_text = parsed["response"]
        except Exception:
            pass

        # Xóa các file tạm giải mã từ base64 sau khi hoàn thành
        for path in created_temp_files:
            try:
                if os.path.exists(path):
                    os.remove(path)
            except Exception:
                pass
                
        image_base64 = None
        video_base64 = None
        img_path = target_path
        
        match = re.search(r"(?:IMAGE_PATH|VIDEO_PATH)=([^\s]+)", reply_text)
        if match:
            img_path = match.group(1).strip()
            if img_path.endswith(']'): img_path = img_path[:-1]

        if not os.path.exists(img_path) and image_uid:
            import glob
            brain_dir = os.path.expanduser("~/.gemini/antigravity-cli/brain")
            pattern_png = os.path.join(brain_dir, "**", f"*{image_uid}*.png")
            pattern_jpg = os.path.join(brain_dir, "**", f"*{image_uid}*.jpg")
            pattern_jpeg = os.path.join(brain_dir, "**", f"*{image_uid}*.jpeg")
            found_files = glob.glob(pattern_png, recursive=True) + glob.glob(pattern_jpg, recursive=True) + glob.glob(pattern_jpeg, recursive=True)
            if found_files: img_path = found_files[0]

        if os.path.exists(img_path):
            if img_path.endswith('.png') or img_path.endswith('.jpg') or img_path.endswith('.jpeg'):
                try:
                    img = Image.open(img_path)
                    width, height = img.size
                    target_width = 2048  # Mặc định tối thiểu 2K HD cho mọi ảnh sinh ra
                    if "4k" in full_prompt.lower(): 
                        target_width = 3840
                    elif "2k" in full_prompt.lower(): 
                        target_width = 2560
                    
                    if width < target_width:
                        aspect = height / width
                        new_height = int(target_width * aspect)
                        resample_method = getattr(Image, "Resampling", Image).LANCZOS
                        img = img.resize((target_width, new_height), resample_method)
                        img.save(img_path, format="PNG", quality=100, optimize=True)
                except Exception as e:
                    print("Lỗi khi upscale/tối ưu chất lượng ảnh:", e)
                    pass
                
            with open(img_path, "rb") as f:
                b64_data = base64.b64encode(f.read()).decode("utf-8")
                if img_path.endswith('.mp4'): video_base64 = b64_data
                else: image_base64 = b64_data
            try: os.remove(img_path)
            except: pass
        
        if image_base64 or video_base64:
            if image_base64:
                data_uri = f"data:image/png;base64,{image_base64}"
                img_md = f"\n\n![Ảnh 3D vừa tạo]({data_uri})\n\n"
            else:
                data_uri = f"data:video/mp4;base64,{video_base64}"
                img_md = f"\n\n<video src=\"{data_uri}\" controls></video>\n\n"

            final_chunk = {
                "success": True,
                "content": img_md,
                "image": data_uri if image_base64 else None,
                "url": data_uri,
                "image_url": data_uri if image_base64 else None,
                "image_base64": image_base64,
                "video_base64": video_base64,
                "conversation_id": conversation_id
            }
            await yield_callback(f"data: {json.dumps(final_chunk)}\n\n")
        else:
            final_chunk = {"success": True, "content": "", "conversation_id": conversation_id}
            await yield_callback(f"data: {json.dumps(final_chunk)}\n\n")
            
    except Exception as e:
        err_chunk = {"success": False, "error": str(e)}
        await yield_callback(f"data: {json.dumps(err_chunk)}\n\n")


async def extract_request_data(request: Request, default_sys_instruction: str = "Bạn là trợ lý AI.", default_is_automation: bool = False) -> dict:
    content_type = request.headers.get("content-type", "")
    file_paths = []
    files_data = []
    
    if "application/json" in content_type:
        body = await request.json()
        conversation_id = body.get("conversation_id") or str(uuid.uuid4())
        
        # 1. Parse files_data base64
        raw_files_data = body.get("files_data", [])
        if isinstance(raw_files_data, list):
            for item in raw_files_data:
                if isinstance(item, dict):
                    if "filename" in item and "content_base64" in item:
                        files_data.append(item)
                    elif "path" in item:
                        file_paths.append(item["path"])

        # 2. Parse file_paths / files_path / file_path / reference_files / files
        for key in ("file_paths", "files_path", "file_path", "reference_files", "files"):
            raw_val = body.get(key)
            if isinstance(raw_val, str) and raw_val.strip():
                file_paths.append(raw_val.strip())
            elif isinstance(raw_val, list):
                for item in raw_val:
                    if isinstance(item, str) and item.strip():
                        file_paths.append(item.strip())
                    elif isinstance(item, dict) and "path" in item and item["path"]:
                        file_paths.append(item["path"])

        prompt = body.get("prompt", "Xin chào" if not default_is_automation else "")
        system_instructions = body.get("system_instructions", default_sys_instruction)
        model = body.get("model")
        gemini_api_key = body.get("gemini_api_key")
        tts_voice = body.get("tts_voice", "vi-VN-HoaiMyNeural")
        tts_rate = body.get("tts_rate", "+0%")
        is_auto = body.get("is_automation", default_is_automation)
    else:
        form = await request.form()
        conversation_id = form.get("conversation_id") or str(uuid.uuid4())
        
        files = form.getlist("files")
        for file in files:
            if hasattr(file, "filename") and file.filename:
                content = await file.read()
                files_data.append({
                    "filename": file.filename,
                    "content_base64": base64.b64encode(content).decode('utf-8')
                })
            elif isinstance(file, str) and file.strip():
                file_paths.append(file.strip())

        for key in ("file_paths", "files_path", "file_path", "reference_files"):
            for val in form.getlist(key):
                if isinstance(val, str) and val.strip():
                    file_paths.append(val.strip())

        prompt = form.get("prompt", "Xin chào" if not default_is_automation else "")
        system_instructions = form.get("system_instructions", default_sys_instruction)
        model = form.get("model")
        gemini_api_key = form.get("gemini_api_key")
        tts_voice = form.get("tts_voice", "vi-VN-HoaiMyNeural")
        tts_rate = form.get("tts_rate", "+0%")
        is_auto = default_is_automation or bool(form.get("is_automation"))

    return {
        "prompt": prompt,
        "system_instructions": system_instructions,
        "model": model,
        "gemini_api_key": gemini_api_key,
        "conversation_id": conversation_id,
        "tts_voice": tts_voice,
        "tts_rate": tts_rate,
        "files_data": files_data,
        "file_paths": file_paths,
        "is_automation": is_auto
    }

def handle_agent_streaming(req_data: dict):
    if connected_pcs:
        # Chuyển tiếp (Relay) qua WebSocket tới PC đầu tiên đang kết nối
        task_id = str(uuid.uuid4())
        pc_ws = connected_pcs[0]
        req_data["task_id"] = task_id
        active_tasks[task_id] = asyncio.Queue()
        
        async def send_to_ws():
            await pc_ws.send_json(req_data)
        asyncio.create_task(send_to_ws())
        
        async def relay_stream():
            try:
                while True:
                    try:
                        data = await asyncio.wait_for(active_tasks[task_id].get(), timeout=15.0)
                        if data.get("done"):
                            break
                        if "chunk" in data:
                            yield data["chunk"]
                    except asyncio.TimeoutError:
                        # Gửi ping SSE comment định kỳ mỗi 15s để duy trì kết nối HTTP với Cloudflare (chống lỗi 524)
                        yield ": ping\n\n"
            finally:
                if task_id in active_tasks:
                    del active_tasks[task_id]
        
        return StreamingResponse(relay_stream(), media_type="text/event-stream")
    else:
        # Chạy trực tiếp (Local execution)
        async def generate_stream_local():
            queue_local = asyncio.Queue()
            async def local_yield(chunk):
                await queue_local.put(chunk)
            
            async def run_task_in_bg():
                await run_agent_task_logic(req_data, local_yield)
                await queue_local.put(None)
                
            asyncio.create_task(run_task_in_bg())
            
            while True:
                try:
                    chunk = await asyncio.wait_for(queue_local.get(), timeout=15.0)
                    if chunk is None:
                        break
                    yield chunk
                except asyncio.TimeoutError:
                    # Gửi ping SSE comment định kỳ mỗi 15s để duy trì kết nối HTTP với Cloudflare (chống lỗi 524)
                    yield ": ping\n\n"

        return StreamingResponse(generate_stream_local(), media_type="text/event-stream")

@app.post("/api/chat")
async def chat_endpoint(request: Request):
    """
    API Chuyên biệt cho Trò chuyện & Sinh nội dung AI tổng quát (General AI Chat).
    """
    x_api_key = request.headers.get("x-api-key")
    if x_api_key != SECRET_AGENT_KEY:
        raise HTTPException(status_code=401, detail="Unauthorized: Sai hoặc thiếu API Key (x-api-key)")
    try:
        req_data = await extract_request_data(request, default_sys_instruction="Bạn là trợ lý AI.", default_is_automation=False)
        return handle_agent_streaming(req_data)
    except Exception as e:
        return JSONResponse(status_code=500, content={"success": False, "result": "", "error": str(e)})

@app.post("/api/automation")
@app.post("/api/automate")
async def automation_endpoint(request: Request):
    """
    API Chuyên biệt cho Tự động hóa hệ thống máy tính (OS Automation Task).
    Tách biệt hoàn toàn với /api/chat để thực thi các tác vụ điều khiển desktop và hệ thống.
    """
    x_api_key = request.headers.get("x-api-key")
    if x_api_key != SECRET_AGENT_KEY:
        raise HTTPException(status_code=401, detail="Unauthorized: Sai hoặc thiếu API Key (x-api-key)")
    try:
        req_data = await extract_request_data(request, default_sys_instruction="Bạn là AI Agent tự động hóa hệ thống máy tính.", default_is_automation=True)
        return handle_agent_streaming(req_data)
    except Exception as e:
        return JSONResponse(status_code=500, content={"success": False, "result": "", "error": str(e)})


# ============================================================
# PROFILE SCENE DATA API — /api/profile/scene/{username}
# Quy tắc Bắt buộc: CHỈ xử lý duy nhất các nghiệp vụ 3D PlayCanvas
# (CRUD profile.json & assets GLB/textures) có trong Documents/ai.type/data/profiles.
# TUYỆT ĐỐI KHÔNG xử lý bất kỳ công việc nào khác (KHÔNG LLM, KHÔNG chat, KHÔNG OS commands).
# Storage: /home/yenai/Documents/ai.type/data/profiles/{username}/
# ============================================================
PROFILES_ROOT = os.path.expanduser(os.path.join("~", "Documents", "ai.type", "data", "profiles"))

def _profile_dir(username: str) -> str:
    """Trả về đường dẫn thư mục profile của user, tạo nếu chưa có."""
    # Sanitize username: chỉ cho phép ký tự an toàn
    safe = re.sub(r"[^a-zA-Z0-9_\-\.]", "_", username)
    if not safe or safe in (".", ".."):
        raise HTTPException(status_code=400, detail="Username không hợp lệ")
    path = os.path.join(PROFILES_ROOT, safe)
    os.makedirs(path, exist_ok=True)
    return path

def _profile_json_path(username: str) -> str:
    return os.path.join(_profile_dir(username), "profile.json")

DEFAULT_PROFILE = {
    "id": "profile_admin",
    "username": "",
    "character": {
        "id": "char_main",
        "name": "Yên",
        "animationState": "Walk",
        "skin": "default",
        "glbUrl": None
    },
    "scene": {
        "id": "room_306",
        "cameraPreset": "DESK_VIEW",
        "lightsOn": True
    },
    "customizations": {},
    "createdAt": None,
    "updatedAt": None
}

import datetime

@app.get("/api/profile/scene/{username}")
async def get_profile_scene(username: str, request: Request):
    """Lấy profile data của user (character, scene settings, customizations)."""
    x_api_key = request.headers.get("x-api-key")
    if x_api_key != SECRET_AGENT_KEY:
        raise HTTPException(status_code=401, detail="Unauthorized")
    
    profile_path = _profile_json_path(username)
    if not os.path.exists(profile_path):
        # Tạo profile mặc định cho user mới
        profile = {**DEFAULT_PROFILE, "username": username, "createdAt": datetime.datetime.utcnow().isoformat()}
        with open(profile_path, "w", encoding="utf-8") as f:
            json.dump(profile, f, ensure_ascii=False, indent=2)
        return {"success": True, "data": profile, "isNew": True}
    
    with open(profile_path, "r", encoding="utf-8") as f:
        profile = json.load(f)
    
    # Liệt kê các asset files của user (GLB, textures...)
    asset_dir = os.path.join(_profile_dir(username), "assets")
    assets = []
    if os.path.exists(asset_dir):
        for fname in os.listdir(asset_dir):
            fpath = os.path.join(asset_dir, fname)
            if os.path.isfile(fpath):
                assets.append({
                    "name": fname,
                    "size": os.path.getsize(fpath),
                    "url": f"/api/profile/scene/assets/{username}/{fname}"
                })
    profile["_assets"] = assets
    
    return {"success": True, "data": profile}

@app.put("/api/profile/scene/{username}")
async def put_profile_scene(username: str, request: Request):
    """Lưu và đồng bộ profile data 3D của user vào profile.json và các tệp cấu hình 3D phụ."""
    x_api_key = request.headers.get("x-api-key")
    if x_api_key != SECRET_AGENT_KEY:
        raise HTTPException(status_code=401, detail="Unauthorized")
    
    body = await request.json()
    updated_profile = update_playcanvas_3d_profile(username, body)
    return {"success": True, "data": updated_profile}

@app.post("/api/profile/scene/assets/{username}/{filename}")
async def upload_profile_asset(username: str, filename: str, request: Request, file: UploadFile = File(...)):
    """Upload asset file (GLB character, texture...) vào thư mục profile của user."""
    x_api_key = request.headers.get("x-api-key")
    if x_api_key != SECRET_AGENT_KEY:
        raise HTTPException(status_code=401, detail="Unauthorized")
    
    # Chỉ cho phép một số extension an toàn
    allowed_exts = {".glb", ".gltf", ".png", ".jpg", ".jpeg", ".webp", ".json"}
    ext = os.path.splitext(filename)[1].lower()
    if ext not in allowed_exts:
        raise HTTPException(status_code=400, detail=f"File type not allowed: {ext}")
    
    # Sanitize filename
    safe_name = re.sub(r"[^a-zA-Z0-9_\-\.]", "_", filename)
    asset_dir = os.path.join(_profile_dir(username), "assets")
    os.makedirs(asset_dir, exist_ok=True)
    
    file_path = os.path.join(asset_dir, safe_name)
    contents = await file.read()
    with open(file_path, "wb") as f:
        f.write(contents)
    
    return {
        "success": True,
        "filename": safe_name,
        "size": len(contents),
        "url": f"/api/profile/scene/assets/{username}/{safe_name}"
    }

@app.get("/api/profile/scene/assets/{username}/{filename}")
async def get_profile_asset(username: str, filename: str, request: Request):
    """Lấy asset file của user (GLB, texture...)."""
    # Asset files là public (không cần auth) để PlayCanvas có thể load trực tiếp
    safe_name = re.sub(r"[^a-zA-Z0-9_\-\.]", "_", filename)
    asset_dir = os.path.join(_profile_dir(username), "assets")
    file_path = os.path.join(asset_dir, safe_name)
    
    if not os.path.isfile(file_path):
        raise HTTPException(status_code=404, detail="Asset not found")
    
    return FileResponse(file_path)


@app.get("/api/shutdown")
def shutdown_endpoint():
    import threading
    def die(): os._exit(0)
    threading.Timer(0.5, die).start()
    return {"success": True, "message": "Shutting down"}


async def relay_client_loop(url: str):
    print(f"Bắt đầu chạy chế độ WebSocket Relay Client. Đang kết nối tới: {url}")
    while True:
        try:
            ws_url = url + f"?api_key={SECRET_AGENT_KEY}"
            async with websockets.connect(ws_url) as ws:
                print("Đã kết nối thành công tới máy chủ Relay!")
                while True:
                    req_data = await ws.recv()
                    req_data = json.loads(req_data)
                    task_id = req_data.get("task_id")
                    if not task_id:
                        continue
                        
                    print(f"Nhận được yêu cầu xử lý từ điện thoại, Task ID: {task_id}")
                    
                    async def send_chunk_to_ws(chunk: str):
                        await ws.send(json.dumps({"task_id": task_id, "chunk": chunk}))
                        
                    async def process_and_report():
                        try:
                            await run_agent_task_logic(req_data, send_chunk_to_ws)
                        finally:
                            await ws.send(json.dumps({"task_id": task_id, "done": True}))
                            
                    asyncio.create_task(process_and_report())
        except Exception as e:
            print("Mất kết nối với Relay Server, thử lại sau 3 giây...", str(e))
            await asyncio.sleep(3)


@app.on_event("startup")
async def startup_event():
    if RELAY_TO:
        asyncio.create_task(relay_client_loop(RELAY_TO))

import sys
if getattr(sys, 'frozen', False):
    base_dir = sys._MEIPASS
else:
    base_dir = os.path.dirname(os.path.abspath(__file__))
DIST_DIR = os.path.join(base_dir, "sontinh_dist")
if os.path.exists(DIST_DIR):
    assets_dir = os.path.join(DIST_DIR, "assets")
    if os.path.exists(assets_dir):
        app.mount("/assets", StaticFiles(directory=assets_dir), name="assets")

    @app.get("/")
    async def serve_index():
        return FileResponse(os.path.join(DIST_DIR, "index.html"))

    @app.get("/{full_path:path}")
    async def serve_static_or_spa(request: Request, full_path: str):
        if full_path.startswith("api/") or full_path.startswith("ws/"):
            raise HTTPException(status_code=404, detail="Not Found")
        file_path = os.path.join(DIST_DIR, full_path)
        if os.path.isfile(file_path):
            return FileResponse(file_path)
        return FileResponse(os.path.join(DIST_DIR, "index.html"))

if __name__ == "__main__":
    uvicorn.run(app, host="127.0.0.1", port=54321)
